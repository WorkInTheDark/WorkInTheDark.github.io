#!/usr/bin/env python3
"""The Chinese translation as a spreadsheet Arthur can edit, and back.

  export: ruby tools/phd-sim/export-data.rb
          node --experimental-default-type=module tools/phd-sim/i18n-review.mjs <rows.json> --json
          python3 tools/phd-sim/i18n-xlsx.py export <rows.json> <out.xlsx> [--literal <git ref>]
                  [--prev <git ref>] [--reviewed <reviewed.json>] [--notes <notes.json>] [--pending <git ref> --round <n>]
  import: python3 tools/phd-sim/i18n-xlsx.py import <edited.xlsx>

For a later round, --prev marks what changed since that version (with the old text beside it),
--reviewed marks rows Arthur already reviewed ({"reviewed": [[table, English], ...]}), and --notes
explains changes to them ({"<table>\t<English>": "why"}). --pending <ref> keeps only the lines that are new
or changed since that version (the last round Arthur approved), each marked new or changed, with the
earlier Chinese beside a changed line; --round numbers the sheet (round 4 follows round 3's approval).

Rows are matched by table + English, so sorting or filtering the sheet is fine. An edited 中文 cell
replaces that line's translation in _data/phd_sim_i18n/zh/<table>.yml, and in any other table with
the same English (the game shares one translation per English line). Notes are printed for Claude.
"""
import json, os, re, subprocess, sys
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
ZH_DIR = os.path.join(ROOT, "_data", "phd_sim_i18n", "zh")
SHEETS = [("事件 Events", lambda s: s.startswith("events/")), ("游戏文本 Game text", lambda s: not s.startswith("events/") and s != "code"),
          ("界面 Interface", lambda s: s == "code")]
WIDTH = {"新增/改动 Status": 12, "表 Table": 16, "位置 Where": 30, "English": 60, "中文（在此修改）": 60, "本次改写 Localized": 10, "本轮改动 Changed": 10,
         "Claude 的说明": 36, "上一版 Previous": 46, "已审 Reviewed": 9, "直译版 Literal first pass": 40, "备注 Notes for Claude": 30}
GREY = ("表 Table", "位置 Where", "上一版 Previous", "直译版 Literal first pass")
WRAP = ("English", "中文（在此修改）", "Claude 的说明", "上一版 Previous", "直译版 Literal first pass", "备注 Notes for Claude")
FONT = "Arial"
PAIR = re.compile(r'^- en: (".*")\n  zh: (".*")$', re.M)


def literal_tables(ref):
    """The first, literal translation, read from git, as {table: {en: zh}}."""
    out = {}
    if not ref:
        return out
    files = subprocess.run(["git", "ls-tree", "-r", "--name-only", ref, "_data/phd_sim_i18n/zh"], cwd=ROOT, capture_output=True, text=True).stdout.split()
    for f in files:
        if not f.endswith(".yml") or f.endswith("_stale.yml"):
            continue
        src = f[len("_data/phd_sim_i18n/zh/"):-4]
        text = subprocess.run(["git", "show", ref + ":" + f], cwd=ROOT, capture_output=True, text=True).stdout
        out[src] = {r["en"]: r.get("zh", "") for r in (yaml.safe_load(text) or [])}
    return out


def export(rows_json, out_path, ref, prev_ref=None, reviewed_json=None, notes_json=None, pending_ref=None, round_n=3):
    data = json.load(open(rows_json, encoding="utf-8"))
    if pending_ref:
        return export_pending(data, out_path, pending_ref, round_n)
    lit = literal_tables(ref)
    prev = literal_tables(prev_ref) if prev_ref else None
    reviewed = {tuple(r) for r in json.load(open(reviewed_json, encoding="utf-8"))["reviewed"]} if reviewed_json else None
    notes = json.load(open(notes_json, encoding="utf-8")) if notes_json else {}
    head = ["表 Table", "位置 Where", "English", "中文（在此修改）"]
    head += ["本轮改动 Changed", "Claude 的说明", "上一版 Previous"] if prev is not None else ["本次改写 Localized"]
    head += ["已审 Reviewed"] if reviewed is not None else []
    head += ["直译版 Literal first pass", "备注 Notes for Claude"]
    wb = Workbook()
    readme = wb.active
    readme.title = "说明 README"
    lines = [
        ("读博模拟器 · 中文翻译表", True),
        ("每一行是游戏里的一句话。只改黄色的「中文（在此修改）」一列；其他列请不要改（导入时按「表 + English」对行）。", False),
        ("可以排序、筛选、加批注；想让 Claude 处理的问题写在「备注 Notes for Claude」里。", False),
        ("{advisor}、{labmate}、{n} 这类花括号是游戏自动填的内容（人名、数字等），请原样保留，可以挪位置。", False),
        ("「直译版」是第一版直译，只供对照；「本次改写」标 ✓ 的行是这次本地化改过的。", False) if prev is None else
        ("「已审」标 ✓ 的是你审过的行。筛选「本轮改动」= ✓ 能看到 Claude 这一轮改了哪些；改你审过的行时，理由写在「Claude 的说明」。「上一版」是这一轮之前的中文，「直译版」是最早的直译。", False),
        ("同一句英文在多个表里出现时，游戏只用一种中文：改其中一处，导入时会同步到所有地方。", False),
        ("改完把文件发回给 Claude（或放回原处），Claude 会导入、检查并在游戏里测试。", False),
        ("", False),
        ("Each row is one line of the game. Edit only the yellow 中文 column; rows are matched by table + English, so sorting and filtering are safe.", False),
        ("Keep {tokens} like {advisor} or {n} exactly (you can move them). Put anything you want Claude to handle in the Notes column.", False),
        ("", False),
        ("示例 Example", True),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = readme.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT, bold=bold, size=14 if i == 1 else 11)
    ex_row = len(lines) + 1
    for j, h in enumerate(["English", "原来的中文", "你改成"], 1):
        readme.cell(row=ex_row, column=j, value=h).font = Font(name=FONT, bold=True)
    for j, v in enumerate(["{advisor} says \"fine\" in a way that is not quite fine.", "{advisor}说了句“行吧”，听起来不太行。", "{advisor}回了句“行吧”。懂的都懂，这不太行。"], 1):
        c = readme.cell(row=ex_row + 1, column=j, value=v)
        c.font = Font(name=FONT)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 60
    readme.column_dimensions["B"].width = 44
    readme.column_dimensions["C"].width = 44

    edit_fill = PatternFill("solid", fgColor="FFF6CC")
    head_fill = PatternFill("solid", fgColor="283454")
    thin = Side(style="thin", color="DDDDDD")
    counts = {}
    for title, pick in SHEETS:
        ws = wb.create_sheet(title)
        for j, h in enumerate(head, 1):
            c = ws.cell(row=1, column=j, value=h)
            c.font = Font(name=FONT, bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        r = 2
        for f in data["files"]:
            if not pick(f["src"]):
                continue
            old = lit.get(f["src"], {})
            before = prev.get(f["src"], {}) if prev is not None else {}
            for g in f["groups"]:
                for en, zh in g["rows"]:
                    was = old.get(en, "")
                    v = {"表 Table": f["src"], "位置 Where": g["label"], "English": en, "中文（在此修改）": zh,
                         "本次改写 Localized": "✓" if was and was != zh else "", "直译版 Literal first pass": was}
                    if prev is not None:
                        b = before.get(en, "")
                        v.update({"本轮改动 Changed": "✓" if b and b != zh else "", "上一版 Previous": b if b != zh else "",
                                  "Claude 的说明": notes.get(f["src"] + "\t" + en, "")})
                    if reviewed is not None:
                        v["已审 Reviewed"] = "✓" if (f["src"], en) in reviewed else ""
                    for j, h in enumerate(head, 1):
                        c = ws.cell(row=r, column=j, value=v.get(h, ""))
                        c.font = Font(name=FONT, size=10, color="777777" if h in GREY else "000000")
                        c.alignment = Alignment(wrap_text=h in WRAP, vertical="top")
                        c.border = Border(bottom=thin)
                    ws.cell(row=r, column=4).fill = edit_fill
                    r += 1
        counts[title] = r - 2
        for j, h in enumerate(head, 1):
            ws.column_dimensions[get_column_letter(j)].width = WIDTH[h]
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = "A1:" + get_column_letter(len(head)) + str(max(2, r - 1))
    wb.save(out_path)
    print("wrote", out_path, counts)


def export_pending(data, out_path, ref, round_n=3):
    """Only what Arthur hasn't reviewed: lines added or changed since ref. The game shares one Chinese
    line per English line, so a line counts as new only if its English wasn't in any table at ref."""
    then = {}
    for tbl in literal_tables(ref).values():
        then.update(tbl)
    head = ["新增/改动 Status", "表 Table", "位置 Where", "English", "中文（在此修改）", "上一版 Previous", "备注 Notes for Claude"]
    edit_col = head.index("中文（在此修改）") + 1
    wb = Workbook()
    readme = wb.active
    readme.title = "说明 README"
    edit_fill = PatternFill("solid", fgColor="FFF6CC")
    head_fill = PatternFill("solid", fgColor="283454")
    thin = Side(style="thin", color="DDDDDD")
    counts, total, new_n, changed_n = {}, 0, 0, 0
    sheets = []
    for title, pick in SHEETS:
        ws = wb.create_sheet(title)
        for j, h in enumerate(head, 1):
            c = ws.cell(row=1, column=j, value=h)
            c.font = Font(name=FONT, bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        r = 2
        for f in data["files"]:
            if not pick(f["src"]):
                continue
            for g in f["groups"]:
                for en, zh in g["rows"]:
                    if en in then and then[en] == zh:
                        continue
                    is_new = en not in then
                    new_n += is_new
                    changed_n += not is_new
                    v = {"新增/改动 Status": "新增 New" if is_new else "改动 Changed", "表 Table": f["src"], "位置 Where": g["label"],
                         "English": en, "中文（在此修改）": zh, "上一版 Previous": "" if is_new else then[en]}
                    for j, h in enumerate(head, 1):
                        c = ws.cell(row=r, column=j, value=v.get(h, ""))
                        c.font = Font(name=FONT, size=11 if h in ("English", "中文（在此修改）") else 10,
                                      color="777777" if h in GREY else ("B4532A" if h == "新增/改动 Status" and not is_new else "000000"))
                        c.alignment = Alignment(wrap_text=h in WRAP, vertical="top")
                        c.border = Border(bottom=thin)
                    ws.cell(row=r, column=edit_col).fill = edit_fill
                    r += 1
        counts[title] = r - 2
        total += r - 2
        for j, h in enumerate(head, 1):
            ws.column_dimensions[get_column_letter(j)].width = WIDTH[h]
        ws.freeze_panes = "E2"
        ws.auto_filter.ref = "A1:" + get_column_letter(len(head)) + str(max(2, r - 1))
        sheets.append(ws)
    for ws in [x for x in sheets if x.max_row < 2]:
        wb.remove(ws)
    lines = [
        ("读博模拟器 · 待审中文（第 " + str(round_n) + " 轮）", True),
        ("这里只有你上次确认（第 " + str(round_n - 1) + " 轮）之后新增或改动的中文，共 " + str(total) + " 行：新增 " + str(new_n) + " 行，改动 " + str(changed_n) + " 行。", False),
        ("每个分页按游戏里的位置排好：事件、游戏文本（事件以外的内容，如会议周、周末活动、结局、成就）、界面文字。", False),
        ("只改黄色的「中文（在此修改）」一列；「上一版」是改动前的中文，只供对照。其他列请不要改（导入时按「表 + English」对行，排序、筛选都没关系）。", False),
        ("{advisor}、{labmate}、{mentee}、{n} 这类花括号是游戏自动填的内容（人名、数字等），请原样保留，可以挪位置。{prof} 显示为“Kim教授”，后面不加空格；{stranger} 是开会时认识的人名。", False),
        ("英文里加了引号的话，中文也用同样多、同样顺序的“”：结局卡片“别人对你说过的话”是按引号的位置摘的。", False),
        ("想让 Claude 处理的问题写在「备注 Notes for Claude」里。改完把文件发回给 Claude，Claude 会导入、检查并在游戏里测试。", False),
        ("", False),
        ("Only the Chinese added or changed since round " + str(round_n - 1) + ", which you approved: " + str(total) + " lines (" + str(new_n) + " new, " + str(changed_n) + " changed).", False),
        ("Edit only the yellow 中文 column; keep {tokens} as they are. Rows are matched by table + English, so sorting and filtering are safe.", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = readme.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT, bold=bold, size=14 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 110
    wb.save(out_path)
    print("wrote", out_path, counts, "| new", new_n, "changed", changed_n)


def read_table(src):
    path = os.path.join(ZH_DIR, src + ".yml")
    return path, open(path, encoding="utf-8").read()


def do_import(xlsx):
    wb = load_workbook(xlsx, data_only=True)
    edits, notes, bad = {}, [], []
    for title, _ in SHEETS:
        if title not in wb.sheetnames:
            continue
        ws = wb[title]
        head = [c.value for c in ws[1]]
        ix = {h: i for i, h in enumerate(head)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            src, en, zh = row[ix["表 Table"]], row[ix["English"]], row[ix["中文（在此修改）"]]
            note = row[ix["备注 Notes for Claude"]] if "备注 Notes for Claude" in ix else None
            if not src or en is None:
                continue
            if isinstance(zh, float) and zh.is_integer():
                zh = int(zh)  # Excel turns a typed "2" into a number
            zh = "" if zh is None else str(zh)
            edits[(src, en)] = zh
            if note:
                notes.append((src, en, zh, str(note)))
    # Shared English: the last edit that differs from the current text wins everywhere.
    current, files = {}, {}
    for dirpath, _, names in os.walk(ZH_DIR):
        for n in names:
            if n.endswith(".yml") and n != "_stale.yml":
                p = os.path.join(dirpath, n)
                src = os.path.relpath(p, ZH_DIR)[:-4]
                text = open(p, encoding="utf-8").read()
                files[src] = (p, text)
                for m in PAIR.finditer(text):
                    current[(src, json.loads(m.group(1)))] = json.loads(m.group(2))
    by_en = {}
    for (src, en), zh in edits.items():
        if (src, en) not in current:
            bad.append((src, en))
            continue
        if zh != current[(src, en)]:
            by_en[en] = zh
    changed = 0
    for src, (p, text) in files.items():
        def rep(m):
            nonlocal changed
            en = json.loads(m.group(1))
            if en in by_en and json.loads(m.group(2)) != by_en[en]:
                changed += 1
                return "- en: " + m.group(1) + "\n  zh: " + json.dumps(by_en[en], ensure_ascii=False)
            return m.group(0)
        new = PAIR.sub(rep, text)
        if new != text:
            open(p, "w", encoding="utf-8").write(new)
    print("edited lines:", len(by_en), "| table rows updated:", changed, "| rows not found (English changed since export?):", len(bad))
    for src, en in bad[:20]:
        print("  not found:", src, "|", en[:70])
    if notes:
        print("notes for Claude:")
        for src, en, zh, note in notes:
            print("  [" + src + "]", en[:60], "→", zh[:40], "| note:", note)


if __name__ == "__main__":
    if len(sys.argv) >= 4 and sys.argv[1] == "export":
        opt = lambda k, d=None: sys.argv[sys.argv.index(k) + 1] if k in sys.argv else d
        export(sys.argv[2], sys.argv[3], opt("--literal", "zh-literal-pass"), opt("--prev"), opt("--reviewed"), opt("--notes"), opt("--pending"), int(opt("--round") or 3))
    elif len(sys.argv) == 3 and sys.argv[1] == "import":
        do_import(sys.argv[2])
    else:
        print(__doc__)
        sys.exit(1)
